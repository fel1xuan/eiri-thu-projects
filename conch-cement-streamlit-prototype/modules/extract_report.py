import json
import os
import re
from pathlib import Path

import pandas as pd
import xlrd
from openpyxl import load_workbook
from openpyxl.styles import Alignment


DEFAULT_CELL_MAP = {
    "生料消耗量t": "G7",
    "本日库存": None,
    "燃煤消耗量t": "U34",
    "石灰石混合料t": "N34",
    "铁质材料（铁粉/铜渣）t": "O34",
    "石灰石配料t": "P34",
    "萤石t": "Q34",
    "煤矸石t": "R34",
    "废布料消耗量t": "W34",
}


def _log(logger, message):
    if logger:
        logger(message)
    else:
        print(message)


def load_cell_map(cell_map_path):
    path = Path(cell_map_path)
    if not path.exists():
        return DEFAULT_CELL_MAP.copy()
    with path.open("r", encoding="utf-8") as f:
        data = json.load(f)
    return {key: (value or None) for key, value in data.items()}


def save_cell_map(cell_map_path, cell_map):
    path = Path(cell_map_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    normalized = {str(key).strip(): (str(value).strip().upper() if value not in [None, ""] else None) for key, value in cell_map.items()}
    with path.open("w", encoding="utf-8") as f:
        json.dump(normalized, f, ensure_ascii=False, indent=2)


def infer_year_month_from_path(file_path):
    """
    从文件名或路径中自动识别年份和月份。

    优先识别文件名中的格式：2026年5月1-31日生产综合日报.xls
    如果文件名中识别不到，则尝试从路径中识别：202605。
    """
    file_name = os.path.basename(file_path)
    match = re.search(r"(20\d{2})年(\d{1,2})月", file_name)
    if match:
        return int(match.group(1)), int(match.group(2))

    match = re.search(r"(20\d{2})(0[1-9]|1[0-2])", str(file_path))
    if match:
        return int(match.group(1)), int(match.group(2))

    raise ValueError("无法从文件名或路径中识别年份和月份，请检查生产日报文件路径。")


def cell_to_row_col(cell):
    if not cell:
        raise ValueError("单元格位置为空")

    col_letters = ""
    row_numbers = ""
    for char in str(cell).strip():
        if char.isalpha():
            col_letters += char.upper()
        elif char.isdigit():
            row_numbers += char

    if not col_letters or not row_numbers:
        raise ValueError(f"无效的单元格位置：{cell}")

    col = 0
    for char in col_letters:
        col = col * 26 + (ord(char) - ord("A") + 1)

    return int(row_numbers) - 1, col - 1


def validate_daily_excel_suffix(input_path):
    if input_path.suffix.lower() not in {".xls", ".xlsx"}:
        raise ValueError("仅支持 .xls 或 .xlsx 格式的生产综合日报。")


def read_cell_value(sheet, cell_ref, engine):
    if engine == "openpyxl":
        return sheet[str(cell_ref).strip().upper()].value
    if engine == "xlrd":
        row_idx, col_idx = cell_to_row_col(cell_ref)
        return sheet.cell_value(row_idx, col_idx)
    raise ValueError(f"不支持的 Excel 读取引擎：{engine}")


def run_extract_report(input_file, output_file, cell_map_path=None, cell_map=None, logger=None):
    """
    从生产综合日报 .xls/.xlsx 中提取指定单元格数据，输出为生产日报提取结果。
    """
    input_path = Path(input_file).expanduser()
    output_path = Path(output_file).expanduser()
    if not input_path.exists():
        raise FileNotFoundError(f"生产日报文件不存在：{input_path}")
    validate_daily_excel_suffix(input_path)

    if cell_map is not None:
        selected_cell_map = cell_map
    elif cell_map_path is not None:
        selected_cell_map = load_cell_map(cell_map_path)
    else:
        selected_cell_map = DEFAULT_CELL_MAP.copy()
    output_path.parent.mkdir(parents=True, exist_ok=True)

    _log(logger, "[0号] 开始提取生产日报...")
    year, month = infer_year_month_from_path(str(input_path))
    _log(logger, f"[0号] 已识别年份：{year}，月份：{month}")

    suffix = input_path.suffix.lower()
    if suffix == ".xlsx":
        workbook = load_workbook(input_path, data_only=True)
        sheet_names = workbook.sheetnames
        engine = "openpyxl"
    elif suffix == ".xls":
        workbook = xlrd.open_workbook(str(input_path))
        sheet_names = workbook.sheet_names()
        engine = "xlrd"
    else:
        raise ValueError("仅支持 .xls 或 .xlsx 格式的生产综合日报。")

    day_sheet_names = sorted([name for name in sheet_names if str(name).isdigit()], key=lambda x: int(x))
    if not day_sheet_names:
        raise ValueError("没有找到名称为数字的 sheet，例如 1、2、3，请检查 Excel 工作表名称。")

    _log(logger, f"[0号] 识别到{len(day_sheet_names)}个日期sheet")
    warnings = []
    results = []

    for sheet_name in day_sheet_names:
        if engine == "openpyxl":
            sheet = workbook[str(sheet_name)]
        else:
            sheet = workbook.sheet_by_name(str(sheet_name))

        day = int(sheet_name)
        date_value = pd.Timestamp(year=year, month=month, day=day)
        row_data = {"日期": date_value.strftime("%Y-%m-%d")}

        for column_name, cell_address in selected_cell_map.items():
            if cell_address in [None, ""]:
                row_data[column_name] = None
                continue

            try:
                row_data[column_name] = read_cell_value(sheet, cell_address, engine)
            except (IndexError, KeyError, ValueError) as exc:
                message = f"[0号] 警告：sheet {sheet_name} 中无法读取 {cell_address}，字段 {column_name} 记为空值。{exc}"
                warnings.append(message)
                _log(logger, message)
                row_data[column_name] = None

        results.append(row_data)

    df = pd.DataFrame(results)
    df.to_excel(output_path, index=False)

    wb = load_workbook(output_path)
    ws = wb.active
    ws.column_dimensions["A"].width = 22
    for cell in ws["A"]:
        if cell.row == 1:
            continue
        cell.alignment = Alignment(horizontal="right")

    for col in range(2, ws.max_column + 1):
        col_letter = ws.cell(row=1, column=col).column_letter
        ws.column_dimensions[col_letter].width = 18

    for cell in ws[1]:
        cell.alignment = Alignment(horizontal="center")

    wb.save(output_path)
    _log(logger, f"[0号] 提取完成，结果已保存：{output_path}")

    return {
        "output_path": str(output_path),
        "rows": len(df),
        "year": year,
        "month": month,
        "day_sheets": day_sheet_names,
        "warnings": warnings,
    }
